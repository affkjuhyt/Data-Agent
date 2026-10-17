"""UC-04-08 worker boundary: validate uploaded rows and register committed tables."""

from __future__ import annotations

import csv
import datetime
import io
import json
import os
import re
import uuid
from typing import Any
from urllib.parse import urlparse

import aioboto3
from sqlalchemy import select, text

from libs.dataagent_db import scoped_session
from libs.dataagent_db.models import catalog, ingest
from services.worker_agent.events import Event, Handler

ROLE = "svc_ingest"
CATALOG_ROLE = "svc_catalog"
BUCKET = "da-ingest"
UPLOADED = "file_upload.uploaded"
MAPPED = "file_upload.mapped"
COMMITTED = "file_upload.committed"
WORKSPACE_SCHEMA = "workspace"
# How the catalog names a workspace column type (catalog.column_meta.normalized_type).
_NORMALIZED = {"text": "string", "numeric": "decimal", "boolean": "boolean", "date": "date"}


def _store_config() -> tuple[str, str, str]:
    parsed = urlparse(os.environ.get("MINIO_URL", ""))
    if not parsed.hostname or not parsed.username or not parsed.password:
        raise RuntimeError("MINIO_URL is not configured")
    endpoint = f"{parsed.scheme}://{parsed.hostname}:{parsed.port or 9000}"
    return endpoint, parsed.username, parsed.password


async def _read_object(key: str) -> bytes:
    endpoint, access_key, secret_key = _store_config()
    async with aioboto3.Session().client(
        "s3",
        endpoint_url=endpoint,
        aws_access_key_id=access_key,
        aws_secret_access_key=secret_key,
    ) as client:
        response = await client.get_object(Bucket=BUCKET, Key=key)
        return await response["Body"].read()


def _rows(content: bytes, file_name: str, sheet: str | None) -> list[list[Any]]:
    extension = os.path.splitext(file_name)[1].lower()
    if extension == ".csv":
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    if extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet_name = sheet or workbook.sheetnames[0]
        rows = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
        workbook.close()
        return rows
    if extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        sheet_name = sheet or workbook.sheet_names()[0]
        worksheet = workbook.sheet_by_name(sheet_name)
        rows = [worksheet.row_values(index) for index in range(worksheet.nrows)]
        workbook.release_resources()
        return rows
    raise ValueError("unsupported upload format")


def _matches(value: Any, kind: str) -> bool:
    if value in (None, ""):
        return True
    if kind == "number":
        try:
            float(str(value).replace(",", "."))
        except ValueError:
            return False
    if kind == "boolean":
        return str(value).strip().lower() in {"true", "false", "0", "1", "yes", "no"}
    return True


def _validate(rows: list[list[Any]], column_map: list[dict[str, Any]]) -> tuple[dict[str, Any], float]:
    if not rows:
        raise ValueError("the upload has no readable rows")
    expected = max(len(rows[0]), len(column_map))
    # The first uploaded event can arrive before the user saves a mapping. Treat unmapped
    # columns as strings so validation still produces a result for the review screen.
    effective_map = [*column_map, *({} for _ in range(expected - len(column_map)))]
    invalid = 0
    for row in rows[1:]:
        if len(row) != expected or any(
            not _matches(
                row[index] if index < len(row) else None,
                str(effective_map[index].get("type", "string")),
            )
            for index in range(expected)
        ):
            invalid += 1
    total = max(0, len(rows) - 1)
    ratio = invalid / total if total else 0.0
    return {
        "row_count": total,
        "sample_rows": min(total, 1000),
        "error_rows": invalid,
        "error_ratio": ratio,
        "checked": True,
    }, ratio


def _table_name(file_name: str, upload_id: uuid.UUID) -> str:
    stem = re.sub(r"[^a-z0-9]+", "_", os.path.splitext(file_name)[0].lower()).strip("_")
    return f"upload_{stem[:40] or 'file'}_{upload_id.hex[-8:]}"


def _identifier(value: Any, fallback: str) -> str:
    name = re.sub(r"[^a-zA-Z0-9_]+", "_", str(value or "").strip()).strip("_").lower()
    return name or fallback


def _column_names(header: list[Any], column_map: list[dict[str, Any]]) -> list[str]:
    names: list[str] = []
    for index, value in enumerate(header):
        mapped = column_map[index].get("name") if index < len(column_map) else None
        name = _identifier(mapped or value, f"column_{index + 1}")
        if name in names:
            name = f"{name}_{index + 1}"
        names.append(name)
    return names


def _postgres_type(column_map: list[dict[str, Any]], index: int) -> str:
    kind = str(column_map[index].get("type", "string")) if index < len(column_map) else "string"
    return {"number": "numeric", "percent": "numeric", "boolean": "boolean", "date": "date"}.get(
        kind, "text"
    )


def _value(value: Any, kind: str) -> Any:
    if value in (None, ""):
        return None
    if kind in {"number", "percent"}:
        return float(str(value).replace(",", "."))
    if kind == "boolean":
        return str(value).strip().lower() in {"true", "1", "yes"}
    if kind == "date" and isinstance(value, datetime.datetime):
        return value.date()
    return value


async def _materialize(
    session: Any, table_name: str, rows: list[list[Any]], column_map: list[dict[str, Any]]
) -> list[dict[str, str]]:
    """Writes the workspace table; returns its columns as created (name, PostgreSQL type)."""
    if not rows:
        raise ValueError("the upload has no readable rows")
    names = _column_names(rows[0], column_map)
    columns = [
        {"name": name, "type": _postgres_type(column_map, index)}
        for index, name in enumerate(names)
    ]
    rows_json = [
        [_value(row[index] if index < len(row) else None, str(column_map[index].get("type", "string")) if index < len(column_map) else "string") for index in range(len(names))]
        for row in rows[1:]
    ]
    for start in range(0, len(rows_json), 1_000):
        await session.execute(
            text("SELECT public.da_materialize_upload(:table_name, CAST(:columns AS jsonb), CAST(:rows AS jsonb), :replace)"),
            {
                "table_name": table_name,
                "columns": json.dumps(columns),
                "rows": json.dumps(rows_json[start : start + 1_000], default=_json_value),
                "replace": start == 0,
            },
        )
    return columns


def _json_value(value: Any) -> str:
    if isinstance(value, datetime.datetime | datetime.date):
        return value.isoformat()
    return str(value)


async def _validate_upload(event: Event) -> None:
    async with scoped_session(ROLE, event.org_id, event.payload.get("uploader_id")) as session:
        row = await session.scalar(select(ingest.FileUpload).where(ingest.FileUpload.id == event.aggregate_id))
        if row is None:
            return
        content = await _read_object(row.object_key)
        validation, ratio = _validate(_rows(content, row.file_name, row.sheet), row.column_map or [])
        row.validation = validation
        row.error_ratio = ratio
        if row.status != "committed":
            row.status = "mapped" if event.event_type == MAPPED else "analyzed"


async def _register_dataset(event: Event) -> None:
    async with scoped_session(ROLE, event.org_id, event.payload.get("uploader_id")) as session:
        row = await session.scalar(select(ingest.FileUpload).where(ingest.FileUpload.id == event.aggregate_id))
        if row is None or row.target_dataset_id is None:
            return
        dataset_id = row.target_dataset_id
        table_name = _table_name(row.file_name, row.id)
        data_space_id = row.data_space_id
        expires_at = row.expires_at
        row_count = (row.validation or {}).get("row_count", 0)
        quality_state = "ok" if (row.error_ratio or 0) == 0 else "warning"
        content = await _read_object(row.object_key)
        upload_rows = _rows(content, row.file_name, row.sheet)
        columns = await _materialize(session, table_name, upload_rows, row.column_map or [])
        row.validation = {**(row.validation or {}), "materialized": True, "table_name": table_name}
    async with scoped_session(CATALOG_ROLE, event.org_id, event.payload.get("uploader_id")) as session:
        if await session.get(catalog.Dataset, dataset_id) is None:
            session.add(
                catalog.Dataset(
                    id=dataset_id,
                    org_id=event.org_id,
                    source_id=None,
                    data_space_id=data_space_id,
                    origin="upload",
                    schema_name="workspace",
                    table_name=table_name,
                    object_kind="table",
                    display_name=row.file_name,
                    row_estimate=row_count,
                    enabled=True,
                    scope="data_space",
                    expires_at=expires_at,
                    quality_state=quality_state,
                    created_by=row.uploader_id,
                    updated_by=row.uploader_id,
                )
            )
        # The columns as materialized, so the semantic layer knows a text month from a date.
        if (
            await session.scalar(
                select(catalog.ColumnMeta.id).where(catalog.ColumnMeta.dataset_id == dataset_id).limit(1)
            )
            is None
        ):
            await session.flush()
            session.add_all(
                catalog.ColumnMeta(
                    id=uuid.uuid4(),
                    org_id=event.org_id,
                    dataset_id=dataset_id,
                    dataset_version=1,
                    name=column["name"],
                    ordinal=ordinal,
                    data_type=column["type"],
                    normalized_type=_NORMALIZED[column["type"]],
                    created_by=row.uploader_id,
                    updated_by=row.uploader_id,
                )
                for ordinal, column in enumerate(columns)
            )


async def handle(event: Event) -> None:
    if event.event_type == COMMITTED:
        await _register_dataset(event)
    else:
        await _validate_upload(event)


HANDLERS: dict[str, Handler] = {UPLOADED: handle, MAPPED: handle, COMMITTED: handle}
