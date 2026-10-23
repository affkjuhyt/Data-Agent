"""UC-01-06: worker-agent writes an answer's export to the object store, and an export that needs
approval opens its HITL request first and runs once a steward approves it."""

from __future__ import annotations

import csv
import io
import json
import uuid
from typing import Any

import pytest

from services.worker_agent import exports, resolutions, storage
from services.worker_agent.events import OutboxTableSource
from services.worker_agent.main import run_once

from . import test_answers
from .conftest import Seed
from .test_answers import Bank, answer_of, build, drain

warehouse = test_answers.warehouse

pytestmark = [pytest.mark.integration]


class MemoryStore:
    def __init__(self) -> None:
        self.files: dict[tuple[str, str], tuple[bytes, str]] = {}

    async def put(self, bucket: str, key: str, body: bytes, content_type: str) -> None:
        self.files[(bucket, key)] = (body, content_type)


@pytest.fixture
def store(monkeypatch: pytest.MonkeyPatch) -> MemoryStore:
    memory = MemoryStore()
    monkeypatch.setattr(storage, "_store", memory)
    return memory


async def _answer(bank: Bank) -> uuid.UUID:
    question = bank.ask("Tỷ lệ nợ xấu theo vùng tháng 9/2026")
    await drain()
    return answer_of(bank.seed, question)[-1]  # type: ignore[no-any-return]


def _export(bank: Bank, answer: uuid.UUID, fmt: str, *, detail: bool, status: str) -> uuid.UUID:
    job = bank.seed.row(
        "qa.export_job",
        org_id=bank.org,
        answer_id=answer,
        requested_by=bank.user,
        format=fmt,
        include_detail_rows=detail,
        status=status,
    )
    bank.seed.row(
        "qa.outbox",
        org_id=bank.org,
        aggregate_type="export_job",
        aggregate_id=job,
        event_type=exports.EVENT_TYPE,
        payload=json.dumps(
            {
                "export_job_id": str(job),
                "answer_id": str(answer),
                "data_space_id": str(bank.space),
                "requested_by": str(bank.user),
                "format": fmt,
                "include_detail_rows": detail,
                "row_count": 2,
                "status": status,
            }
        ),
    )
    return job


async def _drain_exports() -> None:
    source = OutboxTableSource(schema="qa", event_types=frozenset(exports.HANDLERS))
    while await run_once(source, exports.HANDLERS, batch=50):
        pass


def _job(seed: Seed, job: uuid.UUID) -> tuple[Any, ...]:
    return seed.one(  # type: ignore[return-value]
        "SELECT status, object_key, expires_at > now() + interval '23 hours',"
        " expires_at <= now() + interval '24 hours', hitl_request_id FROM qa.export_job"
        " WHERE id = %s",
        job,
    )


async def test_ac_01_06_01_a_csv_export_lands_in_da_exports_with_its_footer_and_audit(
    seed: Seed, warehouse: dict[str, Any], monkeypatch: pytest.MonkeyPatch, store: MemoryStore
) -> None:
    bank = build(seed, warehouse, monkeypatch)
    job = _export(bank, await _answer(bank), "csv", detail=False, status="queued")
    await _drain_exports()
    status, key, after_23h, within_24h, _ = _job(seed, job)
    assert (status, key, after_23h, within_24h) == ("ready", f"{bank.org}/{job}.csv", True, True)
    body, content_type = store.files[("da-exports", key)]
    assert content_type.startswith("text/csv")
    rows = list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))
    assert rows[0] == ["Mã", "Nhãn", "Giá trị", "Đơn vị"]
    assert [r[0] for r in rows[1:3]] == ["F1", "F2"] and rows[1][2].startswith("10.0")
    footer = {r[0]: r[1] for r in rows if len(r) == 2}
    assert footer["Người xuất"] == str(bank.user)
    assert footer["Định nghĩa chỉ số"].startswith("Tỷ lệ nợ xấu: no_xau / du_no * 100")
    assert (
        seed.one(
            "SELECT payload->'after'->>'row_count' FROM qa.outbox"
            " WHERE aggregate_id = %s AND event_type = 'audit.requested'",
            job,
        )
        == "2"
    )


async def test_ac_01_06_02_detail_rows_wait_for_a_steward_then_run_as_the_requester(
    seed: Seed, warehouse: dict[str, Any], monkeypatch: pytest.MonkeyPatch, store: MemoryStore
) -> None:
    bank = build(seed, warehouse, monkeypatch)
    job = _export(bank, await _answer(bank), "xlsx", detail=True, status="waiting_approval")
    await _drain_exports()
    status, key, _, _, request = _job(seed, job)
    assert (status, key) == ("waiting_approval", None) and request is not None
    assert seed.one(
        "SELECT kind, assignee_capability, source_service, context->'excluded_approver_ids'->>0"
        " FROM inbox.hitl_request WHERE id = %s",
        request,
    ) == ("export", "steward", "qa", str(bank.user))
    assert store.files == {}

    steward = seed.row(
        "org.app_user", org_id=bank.org, external_id="kc-st", username="st", display_name="St"
    )
    seed.row(
        "inbox.outbox",
        org_id=bank.org,
        aggregate_type="hitl_request",
        aggregate_id=request,
        event_type=resolutions.EVENT_TYPE,
        payload=json.dumps(
            {
                "hitl_request_id": str(request),
                "source_service": "qa",
                "source_ref": str(job),
                "kind": "export",
                "status": "approved",
                "resolved_by": str(steward),
            }
        ),
    )
    source = OutboxTableSource(schema="inbox", event_types=frozenset(resolutions.HANDLERS))
    while await run_once(source, resolutions.HANDLERS, batch=50):
        pass
    assert _job(seed, job)[0] == "ready"
    from openpyxl import load_workbook

    body, _ = store.files[("da-exports", f"{bank.org}/{job}.xlsx")]
    book = load_workbook(io.BytesIO(body))
    assert book.sheetnames == ["Kết quả", "Chi tiết 1"]
    detail = list(book["Chi tiết 1"].values)
    assert detail[0] == ("vung", "value") and len(detail) == 3
    # The detail rows were read again with the requester's rights, as their own run.
    assert (
        seed.one(
            "SELECT count(*) FROM query.query_run WHERE caller_id = %s AND principal_id = %s",
            job,
            bank.user,
        )
        == 1
    )


async def test_detail_rows_beyond_what_was_approved_are_not_exported(
    seed: Seed, warehouse: dict[str, Any], monkeypatch: pytest.MonkeyPatch, store: MemoryStore
) -> None:
    bank = build(seed, warehouse, monkeypatch)
    job = _export(bank, await _answer(bank), "csv", detail=True, status="queued")
    # The approver saw one row; the query returns two now.
    seed.one("UPDATE qa.export_job SET row_count = 1 WHERE id = %s RETURNING 1", job)
    await _drain_exports()
    assert _job(seed, job)[0] == "failed" and store.files == {}


async def test_a_pdf_export_fails_instead_of_writing_a_wrong_file(
    seed: Seed, warehouse: dict[str, Any], monkeypatch: pytest.MonkeyPatch, store: MemoryStore
) -> None:
    bank = build(seed, warehouse, monkeypatch)
    job = _export(bank, await _answer(bank), "pdf", detail=False, status="queued")
    await _drain_exports()
    assert _job(seed, job)[0] == "failed" and store.files == {}


@pytest.mark.req("AC-01-06-04")
def test_ac_01_06_04_a_text_cell_never_starts_a_formula() -> None:
    assert [exports.safe(v) for v in ("=SUM(A1:A9)", "+84", "-5", "@x", "\t=1", "\r=1", "ok")] == [
        "'=SUM(A1:A9)",
        "'+84",
        "'-5",
        "'@x",
        "'\t=1",
        "'\r=1",
        "ok",
    ]
    # A numeric column's value that is not a plain number is text, and escaped like text.
    assert exports.safe("-2+3+cmd|' /C calc'!A0", numeric=True) == "'-2+3+cmd|' /C calc'!A0"
    # A number of a numeric column is a number, not a formula.
    assert exports.safe("-5.25", numeric=True) == "-5.25"
    assert exports.safe("-5.25", numeric=True, as_number=True) == -5.25
